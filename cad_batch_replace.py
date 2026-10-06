"""通过 AutoCAD COM 批量替换 DWG 中的 Text/MText 内容。

Excel 前两列分别为“原文本”和“替换文本”。该模块只适用于 Windows
上的 AutoCAD/ZWCAD COM 环境，不会在没有 CAD 的环境中运行。
"""

from __future__ import annotations

import argparse
import csv
import logging
import os
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
try:
    from pyautocad import Autocad
except ImportError:  # 允许在非 Windows/CAD 环境检查 Excel 和替换逻辑
    Autocad = None


def load_replacements(excel_path: Path, sheet_name: str | None = None) -> dict[str, str]:
    """读取 Excel 前两列，跳过空原文本和表头。后出现的重复键覆盖前值。"""
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        result: dict[str, str] = {}
        for row_index, row in enumerate(sheet.iter_rows(min_col=1, max_col=2, values_only=True), 1):
            source, target = row
            if row_index == 1 and str(source).strip() in {"原文本", "原文", "旧文本", "source"}:
                continue
            if source is None or str(source) == "":
                continue
            result[str(source)] = "" if target is None else str(target)
        return result
    finally:
        workbook.close()


def replace_text(value: str, replacements: dict[str, str]) -> tuple[str, int]:
    count = 0
    # 先替换较长的键，避免“型号”先替换后破坏“型号A”等更具体规则。
    for source, target in sorted(replacements.items(), key=lambda item: len(item[0]), reverse=True):
        occurrences = value.count(source)
        if occurrences:
            value = value.replace(source, target)
            count += occurrences
    return value, count


def dwg_files(source_dir: Path) -> Iterable[Path]:
    yield from sorted(p for p in source_dir.rglob("*") if p.is_file() and p.suffix.lower() == ".dwg")


def run(source_dir: Path, output_dir: Path, excel_path: Path, sheet_name: str | None = None, progress=None) -> tuple[int, int, int]:
    source_dir = source_dir.resolve()
    output_dir = output_dir.resolve()
    if not source_dir.is_dir():
        raise FileNotFoundError(f"源文件夹不存在：{source_dir}")
    if source_dir == output_dir:
        raise ValueError("输出文件夹不能与源文件夹相同，以免覆盖原图")
    if Autocad is None:
        raise RuntimeError("未安装 pyautocad；请在 Windows 上安装 AutoCAD/ZWCAD 后再运行")
    replacements = load_replacements(excel_path, sheet_name)
    if not replacements:
        raise ValueError("Excel 前两列没有可用的替换规则")
    output_dir.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        filename=output_dir / "cad_batch_replace.log",
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        encoding="utf-8",
    )
    acad = Autocad(create_if_not_exists=True, visible=True)
    processed = changed = failed = 0
    files = [p for p in dwg_files(source_dir) if output_dir not in p.parents]
    for index, source_path in enumerate(files, 1):
        relative = source_path.relative_to(source_dir)
        destination = output_dir / relative
        destination.parent.mkdir(parents=True, exist_ok=True)
        document = None
        try:
            document = acad.app.Documents.Open(str(source_path.resolve()))
            acad.doc = document
            file_changed = 0
            for entity in acad.iter_objects(["Text", "MText"]):
                try:
                    old_value = str(entity.TextString)
                    new_value, count = replace_text(old_value, replacements)
                    if count:
                        entity.TextString = new_value
                        file_changed += count
                except Exception as exc:
                    logging.warning("对象处理失败 %s: %s", source_path, exc)
            document.SaveAs(str(destination.resolve()))
            processed += 1
            changed += file_changed
            logging.info("完成 %s -> %s，替换 %d 处", source_path, destination, file_changed)
        except Exception:
            failed += 1
            logging.exception("处理失败 %s", source_path)
        finally:
            if document is not None:
                try:
                    document.Close(False)
                except Exception:
                    logging.exception("关闭文档失败 %s", source_path)
        if progress:
            progress(index, len(files), source_path.name)
    return processed, changed, failed


def main() -> None:
    parser = argparse.ArgumentParser(description="批量替换 DWG 中的 Text/MText 文本")
    parser.add_argument("source", type=Path, help="包含 DWG 的源文件夹")
    parser.add_argument("excel", type=Path, help="替换表 Excel，前两列为原文本和替换文本")
    parser.add_argument("output", type=Path, help="输出文件夹，不覆盖源图")
    parser.add_argument("--sheet", help="Excel 工作表名称，默认使用当前活动表")
    args = parser.parse_args()
    processed, changed, failed = run(args.source, args.output, args.excel, args.sheet)
    print(f"处理完成：{processed} 个文件，替换 {changed} 处，失败 {failed} 个。")


if __name__ == "__main__":
    main()
